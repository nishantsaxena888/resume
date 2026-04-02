from docx import Document

def add_centene_legacy():
    doc_path = "/Users/nishantsaxena/workspace/resume/update/NISHANT SAXENA___fsg_updated.docx"
    doc = Document(doc_path)
    
    insert_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if "Innovektor Consultancy pvt. Ltd Mumbai" in p.text:
            insert_idx = i
            break
            
    if insert_idx != -1:
        # Insert before Innovektor
        p = doc.paragraphs[insert_idx]
        
        new_p1 = p.insert_paragraph_before("Centene Corporation, St Louis, Missouri\t\t\t\t\t(Jul-2011- Oct 2013)")
        new_p1.style = p.style
        new_p2 = p.insert_paragraph_before("Senior Web Developer")
        new_p2.runs[0].bold = True
        
        new_p3 = p.insert_paragraph_before("Developed and maintained mission-critical healthcare workflows and internal systems.")
        new_p3.style = p.style
        
        new_p4 = p.insert_paragraph_before("")
        
        doc.save(doc_path)
        print("Legacy Centene node injected successfully.")
    else:
        print("Could not find Innovektor!")

if __name__ == "__main__":
    add_centene_legacy()
